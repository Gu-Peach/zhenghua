from __future__ import annotations

import json
from io import BytesIO
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schemas.wire import WireRecord


WIRE_TABLE_HEADERS = [
    "线号",
    "属性",
    "型号",
    "规格",
    "长度",
    "芯",
    "色标",
    "原理号",
    "端子排",
    "起点部位",
    "起点元件",
    "起点名称",
    "起点端子",
    "终点部位",
    "终点元件",
    "终点名称",
    "终点端子",
    "备 注",
    "置信度",
    "来源图片",
    "识别依据",
]

WIRE_TABLE_HEADERS.extend([
    "PDF source pages",
    "source type",
    "external source required",
])

RAW_HEADERS = [
    "source_image",
    "wire_number",
    "attribute",
    "model",
    "spec",
    "length",
    "line_number",
    "core_number",
    "color",
    "terminal_strip",
    "start_part",
    "start_location",
    "start_device",
    "start_name",
    "start_terminal",
    "end_part",
    "end_location",
    "end_device",
    "end_name",
    "end_terminal",
    "remark",
    "confidence",
    "source_note",
    "source_pages",
    "source_type",
    "external_source_required",
    "raw_json",
]


def records_to_xlsx_bytes(
    records: Sequence[WireRecord],
    sheet_title: str = "放线表",
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
) -> bytes:
    workbook = Workbook()
    table_sheet = workbook.active
    table_sheet.title = _safe_sheet_title(sheet_title or "放线表")
    _write_wire_table(table_sheet, records, terminal_strip_mapping=terminal_strip_mapping)

    raw_sheet = workbook.create_sheet("原始提取")
    _write_raw_table(raw_sheet, records)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _write_wire_table(
    sheet: Any,
    records: Sequence[WireRecord],
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
) -> None:
    sheet.append(WIRE_TABLE_HEADERS)
    for record in records:
        data = _record_dict(record)
        row = [
            _first(data, "wire_number", "wire_no", "cable_number", "cable_no"),
            _first(data, "attribute", "property"),
            _first(data, "model", "type", "cable_model"),
            _first(data, "spec", "cable_spec"),
            _first(data, "length", "cable_length"),
            data.get("core_number"),
            data.get("color"),
            data.get("line_number"),
            _terminal_strip_value(data, terminal_strip_mapping),
            _first(data, "start_part", "start_area", "start_section"),
            data.get("start_location") or data.get("start_device"),
            _first(data, "start_name", "start_component_name"),
            _compose_terminal(data.get("start_device"), data.get("start_terminal")),
            _first(data, "end_part", "end_area", "end_section"),
            data.get("end_location") or data.get("end_device"),
            _first(data, "end_name", "end_component_name"),
            _compose_terminal(data.get("end_device"), data.get("end_terminal")),
            data.get("remark"),
            data.get("confidence"),
            data.get("source_image"),
            data.get("source_note"),
            _format_source_pages(data.get("source_pages")),
            data.get("source_type"),
            data.get("external_source_required"),
        ]
        sheet.append([_clean_cell(value) for value in row])
    _style_table(sheet, len(WIRE_TABLE_HEADERS), freeze="A2")


def _write_raw_table(sheet: Any, records: Sequence[WireRecord]) -> None:
    sheet.append(RAW_HEADERS)
    for record in records:
        data = _record_dict(record)
        sheet.append([
            json.dumps(data, ensure_ascii=False) if header == "raw_json"
            else _format_source_pages(data.get(header)) if header == "source_pages"
            else _clean_cell(data.get(header))
            for header in RAW_HEADERS
        ])
    _style_table(sheet, len(RAW_HEADERS), freeze="A2")


def _style_table(sheet: Any, column_count: int, *, freeze: str) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=column_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index in range(1, column_count + 1):
        letter = get_column_letter(index)
        width = _best_width(_iter_column_values(sheet, index), minimum=8, maximum=36)
        sheet.column_dimensions[letter].width = width

    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions


def _iter_column_values(sheet: Any, index: int) -> Iterable[str]:
    for row in sheet.iter_rows(min_col=index, max_col=index, values_only=True):
        value = row[0]
        if value is not None:
            yield str(value)


def _best_width(values: Iterable[str], *, minimum: int, maximum: int) -> int:
    longest = max((len(value) for value in values), default=minimum)
    return max(minimum, min(maximum, longest + 2))


def _record_dict(record: WireRecord) -> dict[str, Any]:
    return record.model_dump(exclude_none=True)


def _first(data: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            return value
    return None


def _clean_cell(value: Any) -> Any:
    if value == "":
        return None
    return value


def _format_source_pages(value: Any) -> str | None:
    if not value:
        return None
    if isinstance(value, (list, tuple, set)):
        pages = sorted({int(page) for page in value if str(page).strip().isdigit()})
        return ",".join(str(page) for page in pages) or None
    return str(value)


def _compose_terminal(device: Any, terminal: Any) -> str | None:
    terminal_text = _to_text(terminal)
    if not terminal_text:
        return None
    if terminal_text in {"*", "PE"} or ":" in terminal_text:
        return terminal_text

    device_text = _to_text(device)
    if not device_text:
        return terminal_text
    normalized_device = device_text.lstrip("-").upper()
    if normalized_device == "XA" or normalized_device == "XH" or normalized_device.startswith("XD"):
        return f"{device_text.lstrip('-')}:{terminal_text}"
    return terminal_text


def _terminal_strip_from_device(device: Any) -> str | None:
    device_text = _to_text(device)
    if not device_text:
        return None
    code = device_text.lstrip("-").upper()
    if code == "XA" or code in {"XD10", "XD11", "XD12"}:
        return "X1"
    if code in {"XD21", "XD23"}:
        return "X21/X23"
    if code in {"XD22", "XD24"}:
        return "X22/X24"
    if code in {"XD3", "XD5"}:
        return "X3"
    if code == "XD4":
        return "X4"
    if code == "XH":
        return "X5"
    return None


def _terminal_strip_value(data: dict[str, Any], mapping: Mapping[str, str] | None) -> str | None:
    device = _to_text(data.get("start_device"))
    if device and mapping:
        normalized = {str(key).lstrip("-").upper(): str(value) for key, value in mapping.items()}
        mapped = normalized.get(device.lstrip("-").upper())
        if mapped:
            return mapped
    return data.get("terminal_strip") or _terminal_strip_from_device(device)


def records_by_segment_to_xlsx_bytes(
    records_by_segment: Mapping[str, Sequence[WireRecord]],
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
) -> bytes:
    """Create one workbook containing one wiring-table/raw pair per segment."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for segment_id, records in records_by_segment.items():
        title = _safe_sheet_title(str(segment_id) or "wiring-table")
        table_sheet = workbook.create_sheet(title)
        _write_wire_table(table_sheet, records, terminal_strip_mapping=terminal_strip_mapping)
        raw_sheet = workbook.create_sheet(_safe_sheet_title(f"{title}-raw"))
        _write_raw_table(raw_sheet, records)

    if not workbook.worksheets:
        sheet = workbook.create_sheet("wiring-table")
        _write_wire_table(sheet, [], terminal_strip_mapping=terminal_strip_mapping)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _safe_sheet_title(title: str) -> str:
    invalid = set('[]:*?/\\')
    cleaned = "".join("_" if char in invalid else char for char in title).strip() or "放线表"
    return cleaned[:31]
