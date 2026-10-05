from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from backend.app.core.config import Settings, default_few_shot_examples_dir, default_prompt_path
from backend.app.services.library_store import list_jobs, make_manifest, manifest_to_job, write_manifest
from backend.app.schemas.library import PageAsset, WireTableGroup
from backend.app.core.env import load_env_files, parse_env_file
from backend.app.schemas.wire import WireRecord
from backend.app.services.excel_writer import records_to_xlsx_bytes
from backend.app.services.pdf_pipeline import _merge_cross_page_groups, _normalize_groups, PageGroup
from backend.app.services.prompt_loader import load_few_shot_examples, load_segment_few_shot_examples
from backend.app.services.test_case_runner import discover_case_dirs, discover_case_sources, run_test_cases
from backend.app.services.vlm_client import (
    ImagePayload,
    VLMError,
    attach_source_metadata,
    build_few_shot_messages,
    build_segment_few_shot_messages,
    parse_vlm_records,
)


def test_default_prompt_path_exists() -> None:
    assert default_prompt_path().is_file()


def test_default_grouping_prompt_path_exists() -> None:
    from backend.app.core.config import default_grouping_prompt_path

    assert default_grouping_prompt_path().is_file()


def test_chat_completions_url_normalizes_provider_root() -> None:
    settings = Settings(
        api_key="key",
        base_url="http://3stooges.chat:4088",
        chat_completions_endpoint=None,
        model="vlm",
        timeout_seconds=120,
        max_pdf_pages=50,
        concurrency=1,
        image_batch_size=4,
        use_response_format=False,
        prompt_path=default_prompt_path(),
    )

    assert settings.chat_completions_url == "http://3stooges.chat:4088/v1/chat/completions"


def test_chat_completions_url_allows_exact_endpoint_override() -> None:
    settings = Settings(
        api_key="key",
        base_url="http://3stooges.chat:4088",
        chat_completions_endpoint="http://3stooges.chat:4088/custom/chat/completions",
        model="vlm",
        timeout_seconds=120,
        max_pdf_pages=50,
        concurrency=1,
        image_batch_size=4,
        use_response_format=False,
        prompt_path=default_prompt_path(),
    )

    assert settings.chat_completions_url == "http://3stooges.chat:4088/custom/chat/completions"


def test_few_shot_example_assets_exist() -> None:
    examples_dir = default_few_shot_examples_dir()
    manifest = json.loads((examples_dir / "manifest.json").read_text(encoding="utf-8"))

    assert len(manifest["cases"]) == 2
    for case in manifest["cases"]:
        assert case["expected_record_count"] > 0
        for image in case["images"]:
            assert (examples_dir / image["path"]).is_file()
        assert (examples_dir / case["id"] / "expected.json").is_file()


def test_multimodal_few_shot_messages_include_images_and_expected_answers() -> None:
    examples = load_few_shot_examples(default_few_shot_examples_dir())
    messages = build_few_shot_messages(examples)

    assert [message["role"] for message in messages] == ["user", "assistant", "user", "assistant"]
    assert all(part["type"] == "image_url" for part in messages[0]["content"][1:])
    assert '"start_terminal": "XD3:3"' in messages[1]["content"]
    assert '"end_terminal": "2-RED"' in messages[3]["content"]


def test_segment_few_shot_examples_include_two_drawing_images_and_merge_answer() -> None:
    examples_dir = default_prompt_path().parent / "example_segment"
    examples = load_segment_few_shot_examples(examples_dir)
    messages = build_segment_few_shot_messages(examples)

    assert len(examples) == 1
    assert len(examples[0].images) == 2
    assert [message["role"] for message in messages] == ["user", "assistant"]
    assert sum(part["type"] == "image_url" for part in messages[0]["content"]) == 2
    assert '"merge": true' in messages[1]["content"]


def test_env_file_parser_and_loader(tmp_path, monkeypatch) -> None:
    env_path = tmp_path / ".env"
    env_path.write_text(
        "# comment\n"
        "VLM_MODEL=case-model\n"
        "VLM_BASE_URL=\"https://example.test/v1\"\n"
        "export VLM_TIMEOUT_SECONDS=30 # seconds\n",
        encoding="utf-8",
    )

    assert parse_env_file(env_path)["VLM_MODEL"] == "case-model"
    assert parse_env_file(env_path)["VLM_BASE_URL"] == "https://example.test/v1"

    monkeypatch.delenv("VLM_MODEL", raising=False)
    load_env_files([env_path])
    assert __import__("os").environ["VLM_MODEL"] == "case-model"


def test_backend_test_case_directories_are_discoverable() -> None:
    case_root = default_prompt_path().parents[2] / "test_case"
    case_dirs = discover_case_dirs(case_root)

    assert [case.name for case in case_dirs] == ["1", "2"]
    assert len(discover_case_sources(case_dirs[0])) == 2
    assert len(discover_case_sources(case_dirs[1])) == 1


def test_test_case_runner_writes_error_file_on_case_failure(tmp_path, monkeypatch) -> None:
    case_dir = tmp_path / "1"
    case_dir.mkdir()
    (case_dir / "drawing.png").write_bytes(b"not-a-real-image-but-good-enough-for-this-test")

    class FakeClient:
        def __init__(self, settings, prompt):
            self.settings = settings
            self.prompt = prompt

        async def extract_images(self, images):
            raise VLMError("simulated timeout")

    def fake_path_to_image_payloads(path, max_pdf_pages):
        from backend.app.services.vlm_client import ImagePayload

        return [ImagePayload(name=path.name, mime_type="image/png", content=path.read_bytes())]

    monkeypatch.setattr("backend.app.services.test_case_runner.VLMClient", FakeClient)
    monkeypatch.setattr("backend.app.services.test_case_runner.path_to_image_payloads", fake_path_to_image_payloads)

    settings = Settings(
        api_key="key",
        base_url="http://example.test",
        chat_completions_endpoint=None,
        model="vlm",
        timeout_seconds=1,
        max_pdf_pages=1,
        concurrency=1,
        image_batch_size=4,
        use_response_format=False,
        prompt_path=default_prompt_path(),
    )

    import asyncio

    results = asyncio.run(run_test_cases(case_root=tmp_path, settings=settings, prompt="prompt"))

    assert results[0].status == "failed"
    assert results[0].error_message == "simulated timeout"
    assert (case_dir / "result.error.json").is_file()
    assert not (case_dir / "result.json").exists()
    assert not (case_dir / "result.xlsx").exists()


def test_parse_vlm_records_accepts_fenced_records_object() -> None:
    records = parse_vlm_records('```json\n{"records":[{"line_number":"003G0121","confidence":"0.9"}]}\n```')

    assert len(records) == 1
    assert records[0].line_number == "003G0121"
    assert records[0].confidence == 0.9


def test_parse_vlm_records_normalizes_start_terminal_and_terminal_strip() -> None:
    records = parse_vlm_records(
        '[{"start_device":"-XD3","start_terminal":"X3:251",'
        '"end_device":"-MA51","end_terminal":"2-RED","line_number":"082G3211"}]'
    )

    assert records[0].start_terminal == "XD3:251"
    assert records[0].end_terminal == "2-RED"
    assert records[0].terminal_strip == "X3"


def test_source_metadata_adds_pdf_pages_and_preserves_external_source() -> None:
    records = [
        WireRecord(source_image="source.pdf#page=2, source.pdf#page=5", source_note="terminal sheet"),
        WireRecord(source_type="external", external_source_required=True, source_note="panel drawing"),
    ]
    images = [
        ImagePayload("source.pdf#page=2", "image/png", b"2", page_number=2),
        ImagePayload("source.pdf#page=5", "image/png", b"5", page_number=5),
    ]

    attach_source_metadata(records, images)

    assert records[0].source_pages == [2, 5]
    assert "PDF" in (records[0].source_note or "")
    assert records[0].source_type == "pdf"
    assert records[1].source_pages == []
    assert records[1].source_type == "external"
    assert records[1].external_source_required is True


def test_records_to_xlsx_bytes_uses_wiring_table_headers() -> None:
    content = records_to_xlsx_bytes(
        [
            WireRecord(
                wire_number="0272",
                attribute="D",
                model="CJV/DA",
                spec="12X1.5",
                length=14,
                core_number=1,
                color="蓝",
                line_number="003G0121",
                terminal_strip="X3",
                start_location="+01F11.3",
                start_device="-XD3",
                start_terminal="3",
                end_location="+01F26",
                end_device="-AIM1",
                end_terminal="X3:3",
                remark="24VDC+",
            )
        ]
    )

    workbook = load_workbook(BytesIO(content))
    sheet = workbook.worksheets[0]

    assert sheet.cell(1, 1).value == "线号"
    assert sheet.cell(1, 8).value == "原理号"
    assert sheet.cell(2, 1).value == "0272"
    assert sheet.cell(2, 8).value == "003G0121"
    assert sheet.cell(2, 13).value == "XD3:3"
    assert sheet.cell(2, 17).value == "X3:3"


def test_page_group_normalization_covers_missing_pages() -> None:
    groups = _normalize_groups(
        [PageGroup(group_id="wire-table-001", title="线表 1", page_numbers=[1, 1, 9])],
        page_count=3,
    )

    assert [group.page_numbers for group in groups] == [[1], [2], [3]]


def test_cross_page_group_merge_uses_embedded_pdf_signatures() -> None:
    images = [
        ImagePayload(name="test.pdf#page=1", mime_type="image/png", content=b"1", page_text="+10F01 082G3211"),
        ImagePayload(name="test.pdf#page=2", mime_type="image/png", content=b"2", blank=True, page_text=""),
        ImagePayload(name="test.pdf#page=3", mime_type="image/png", content=b"3", page_text="+01F11.3 006M0151 =.C+01F26/1.3"),
        ImagePayload(name="test.pdf#page=4", mime_type="image/png", content=b"4", page_text="+01F26 006M0151 +01F11.3-AIM1_X3_3"),
    ]
    groups = _merge_cross_page_groups(
        [
            PageGroup("wire-table-001", "082R50", [1]),
            PageGroup("wire-table-002", "006M01", [3]),
            PageGroup("wire-table-003", "006C01", [4]),
        ],
        images,
    )

    assert [group.page_numbers for group in groups] == [[1], [3, 4]]


def test_cross_page_group_merge_does_not_join_non_adjacent_pages() -> None:
    images = [
        ImagePayload(name="a", mime_type="image/png", content=b"1", page_text="+01F11.3 006M0151"),
        ImagePayload(name="b", mime_type="image/png", content=b"2", page_text="other"),
        ImagePayload(name="c", mime_type="image/png", content=b"3", page_text="+01F11.3 006M0151"),
    ]
    groups = _merge_cross_page_groups(
        [PageGroup("a", "A", [1]), PageGroup("b", "B", [3])],
        images,
    )

    assert [group.page_numbers for group in groups] == [[1], [3]]


def test_library_manifest_round_trip(tmp_path) -> None:
    job_dir = tmp_path / "20260917-demo-abcdef12"
    job_dir.mkdir()
    manifest = make_manifest(
        name="demo",
        source_filename="demo.pdf",
        source_url="/api/v1/library/20260917-demo-abcdef12/source/demo.pdf",
        status="success",
        pages=[PageAsset(page_id="page-1", page_number=1, filename="page_001.png", url="/page")],
        groups=[WireTableGroup(group_id="wire-table-001", title="demo", pages=[1], status="success", record_count=1)],
    )
    write_manifest(job_dir, manifest)

    jobs = list_jobs(tmp_path)
    detail = manifest_to_job(job_dir.name, manifest)

    assert jobs[0].job_id == job_dir.name
    assert detail.record_count == 1


def test_process_pdf_upload_writes_library_outputs(tmp_path, monkeypatch) -> None:
    from fastapi import UploadFile
    from io import BytesIO

    from backend.app.services.vlm_client import ImagePayload
    from backend.app.services.pdf_pipeline import process_pdf_upload

    class FakeClient:
        def __init__(self, settings, prompt):
            self.settings = settings
            self.prompt = prompt

        async def complete(self, *, images, user_text, system_prompt=None, prefix_messages=None):
            return (
                '{"merge":true,"project_no_a":"1002001708",'
                '"project_no_b":"1002001708","drawing_prefix_a":"DQ1002001708",'
                '"drawing_prefix_b":"DQ1002001708","reason":"项目号与图号前缀均一致",'
                '"confidence":0.95,"needs_review":false}'
            )

        async def extract_images(self, images):
            return [WireRecord(wire_number="0272", line_number="003G0121", source_image=images[0].name)]

    def fake_pdf_bytes_to_images(content, name, *, max_pdf_pages, render_scale=2.0):
        return [
            ImagePayload(name=f"{name}#page=1", mime_type="image/png", content=b"page1"),
            ImagePayload(name=f"{name}#page=2", mime_type="image/png", content=b"page2"),
        ]

    monkeypatch.setenv("VLM_MODEL", "vlm")
    monkeypatch.setenv("VLM_API_KEY", "key")
    monkeypatch.setenv("VLM_LIBRARY_ROOT", str(tmp_path))
    monkeypatch.setattr("backend.app.services.pdf_pipeline.VLMClient", FakeClient)
    monkeypatch.setattr("backend.app.agents.wiring_graph.path_to_image_payloads", lambda path, max_pdf_pages, pdf_render_scale=2.0: fake_pdf_bytes_to_images(b"pdf", path.name, max_pdf_pages=max_pdf_pages, render_scale=pdf_render_scale))

    upload = UploadFile(filename="demo.pdf", file=BytesIO(b"pdf"), headers={"content-type": "application/pdf"})
    import asyncio

    response = asyncio.run(process_pdf_upload(file=upload))
    job_dir = tmp_path / response.job.job_id

    assert response.job.status == "success"
    assert response.job.group_count == 1
    assert response.job.record_count == 1
    assert (job_dir / "source" / "demo.pdf").is_file()
    assert (job_dir / "pages" / "page_001.png").is_file()
    assert (job_dir / "groups" / "wire-table-001" / "records.json").is_file()
    assert (job_dir / "groups" / "wire-table-001" / "wiring-table.xlsx").is_file()
