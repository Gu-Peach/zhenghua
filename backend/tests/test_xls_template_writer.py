from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from backend.app.core.terminal_strips import (
    ALLOWED_START_TERMINALS,
    ALLOWED_TERMINAL_STRIPS,
    allowed_start_terminal_code,
    normalize_terminal_strip,
)
from backend.app.schemas.wire import WireRecord
from backend.app.services.xls_template_writer import (
    TEMPLATE_HEADERS,
    load_template_spec,
    records_to_template_xls_bytes,
    records_to_template_xlsx_bytes,
)


TEMPLATE_PATH = Path("case/放线表导入格式(1).xls")


def test_terminal_strip_vocabulary_includes_x0_and_rejects_unknown() -> None:
    assert ALLOWED_TERMINAL_STRIPS == ("X0", "X1", "X21/X23", "X22/X24", "X3", "X4", "X5")
    assert normalize_terminal_strip("XD0") == "X0"
    assert normalize_terminal_strip("XD3") == "X3"
    assert normalize_terminal_strip("anything-else") is None


def test_allowed_start_terminals_are_drawing_codes_not_x_categories() -> None:
    assert "XD21" in ALLOWED_START_TERMINALS
    assert "XA" in ALLOWED_START_TERMINALS
    assert allowed_start_terminal_code("-XD21", "XD21:7") == "XD21"
    assert allowed_start_terminal_code("-K2", "X3:23") is None
    assert allowed_start_terminal_code("X3") is None


def test_real_xls_template_spec_has_23_columns() -> None:
    if not TEMPLATE_PATH.is_file():
        pytest.skip("import template is not present")
    spec = load_template_spec(TEMPLATE_PATH)
    assert spec.sheet_name == "原理图导入表格"
    assert len(spec.headers) == 23
    assert list(spec.headers[: len(TEMPLATE_HEADERS)]) == TEMPLATE_HEADERS
    assert len(spec.column_widths) == 23


def test_native_xls_export_uses_template_columns_and_x0() -> None:
    xlrd = pytest.importorskip("xlrd")
    pytest.importorskip("xlwt")
    content = records_to_template_xls_bytes(
        [
            WireRecord(
                drawing_page="080D01",
                line_number="080D0101",
                start_device="-XD0",
                start_terminal="1",
                end_device="-A1",
                end_terminal="2",
                terminal_strip="invalid-model-value",
                source_pages=[3],
            )
        ],
        template_path=TEMPLATE_PATH,
    )
    workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    sheet = workbook.sheet_by_index(0)
    assert sheet.ncols == 23
    assert sheet.row_values(0)[:12] == TEMPLATE_HEADERS
    assert sheet.cell_value(1, 2) == "X0"
    assert sheet.cell_value(1, 5) == "XD0:1"
    assert all(value in ("", None) for value in sheet.row_values(1)[12:])


def test_xlsx_export_preserves_template_column_count() -> None:
    from openpyxl import load_workbook

    content = records_to_template_xlsx_bytes(
        [WireRecord(start_device="-XD24", start_terminal="5")],
        template_path=TEMPLATE_PATH,
    )
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert sheet.max_column == 23
    assert [sheet.cell(1, column).value for column in range(1, 13)] == TEMPLATE_HEADERS
    assert sheet.cell(2, 3).value == "X22/X24"
