from __future__ import annotations

import json
from collections.abc import Iterable, Mapping, Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agent_service.domain.models.wiring import (
    WireRecord,
)
from agent_service.profiles.export_fields import PREPARED_EXPORT_FIELDS, ExportFields

WIRE_TABLE_HEADERS = [
    "线号",
    "属性",
    "型号",
    "规格",
    "长度",
    "芯",
    "色标",
    "原理号",
    "端子排",
    "起点部位",
    "起点元件",
    "起点名称",
    "起点端子",
    "终点部位",
    "终点元件",
    "终点名称",
    "终点端子",
    "备 注",
    "置信度",
    "来源图片",
    "识别依据",
]

WIRE_TABLE_HEADERS.extend(
    [
        "PDF source pages",
        "source type",
        "external source required",
        "unit_id",
        "connection_id",
        "drawing source pages",
        "status",
        "起点端子排",
        "终点端子排",
        "中间端点",
        "跨页引用",
        "unit identity confidence",
    ]
)

RAW_HEADERS = [
    "source_image",
    "wire_number",
    "attribute",
    "model",
    "spec",
    "length",
    "current",
    "current_basis",
    "current_source_text",
    "line_number",
    "core_number",
    "color",
    "terminal_strip",
    "start_part",
    "start_location",
    "start_device",
    "start_name",
    "start_terminal",
    "end_part",
    "end_location",
    "end_device",
    "end_name",
    "end_terminal",
    "remark",
    "confidence",
    "source_note",
    "source_pages",
    "source_type",
    "external_source_required",
    "unit_id",
    "connection_id",
    "drawing_source_pages",
    "status",
    "start_terminal_strip",
    "end_terminal_strip",
    "intermediate_points",
    "references",
    "unit_identity_confidence",
    "raw_json",
]


def records_to_xlsx_bytes(
    records: Sequence[WireRecord],
    sheet_title: str = "放线表",
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
    export_fields: ExportFields = PREPARED_EXPORT_FIELDS,
) -> bytes:
    workbook = Workbook()
    table_sheet = workbook.active
    table_sheet.title = _safe_sheet_title(sheet_title or "放线表")
    _write_wire_table(
        table_sheet, records, terminal_strip_mapping=terminal_strip_mapping, export_fields=export_fields
    )

    raw_sheet = workbook.create_sheet("原始提取")
    _write_raw_table(raw_sheet, records)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def table_to_xlsx_bytes(
    headers: Sequence[Any],
    rows: Sequence[Sequence[Any]],
    sheet_title: str = "放线表",
) -> bytes:
    """Write the frontend table contract to a single-sheet XLSX workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_title(sheet_title or "放线表")
    sheet.append([_table_cell(value) for value in headers])
    column_count = len(headers)
    for row in rows:
        values = list(row[:column_count])
        if len(values) < column_count:
            values.extend([None] * (column_count - len(values)))
        sheet.append([_table_cell(value) for value in values])
    _style_table(sheet, column_count, freeze="A2")

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _write_wire_table(
    sheet: Any,
    records: Sequence[WireRecord],
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
    export_fields: ExportFields = PREPARED_EXPORT_FIELDS,
) -> None:
    sheet.append(WIRE_TABLE_HEADERS)
    for record in records:
        data = _record_dict(record)
        row = [
            _first(data, "wire_number", "wire_no", "cable_number", "cable_no"),
            _first(data, "attribute", "property"),
            _first(data, "model", "type", "cable_model"),
            _first(data, "spec", "cable_spec"),
            _first(data, "length", "cable_length"),
            data.get("core_number"),
            data.get("color"),
            data.get("line_number"),
            export_fields.terminal_strip(
                data.get("start_device"), data.get("terminal_strip"), terminal_strip_mapping
            ),
            _first(data, "start_part", "start_area", "start_section"),
            data.get("start_location") or data.get("start_device"),
            _first(data, "start_name", "start_component_name"),
            export_fields.terminal(data.get("start_device"), data.get("start_terminal")),
            _first(data, "end_part", "end_area", "end_section"),
            data.get("end_location") or data.get("end_device"),
            _first(data, "end_name", "end_component_name"),
            export_fields.terminal(data.get("end_device"), data.get("end_terminal")),
            data.get("remark"),
            data.get("confidence"),
            data.get("source_image"),
            data.get("source_note"),
            _format_source_pages(data.get("source_pages")),
            data.get("source_type"),
            data.get("external_source_required"),
            data.get("unit_id"),
            data.get("connection_id"),
            _format_string_list(data.get("drawing_source_pages")),
            data.get("status"),
            data.get("start_terminal_strip"),
            data.get("end_terminal_strip"),
            _format_json_cell(data.get("intermediate_points")),
            _format_json_cell(data.get("references")),
            data.get("unit_identity_confidence"),
        ]
        sheet.append([_clean_cell(value) for value in row])
    _style_table(sheet, len(WIRE_TABLE_HEADERS), freeze="A2")


def _write_raw_table(sheet: Any, records: Sequence[WireRecord]) -> None:
    sheet.append(RAW_HEADERS)
    for record in records:
        data = _record_dict(record)
        sheet.append(
            [
                json.dumps(data, ensure_ascii=False)
                if header == "raw_json"
                else _format_source_pages(data.get(header))
                if header == "source_pages"
                else _format_string_list(data.get(header))
                if header == "drawing_source_pages"
                else _format_json_cell(data.get(header))
                if header in {"intermediate_points", "references"}
                else _clean_cell(data.get(header))
                for header in RAW_HEADERS
            ]
        )
    _style_table(sheet, len(RAW_HEADERS), freeze="A2")


def _style_table(sheet: Any, column_count: int, *, freeze: str) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=column_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index in range(1, column_count + 1):
        letter = get_column_letter(index)
        width = _best_width(_iter_column_values(sheet, index), minimum=8, maximum=36)
        sheet.column_dimensions[letter].width = width

    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions


def _iter_column_values(sheet: Any, index: int) -> Iterable[str]:
    for row in sheet.iter_rows(min_col=index, max_col=index, values_only=True):
        value = row[0]
        if value is not None:
            yield str(value)


def _best_width(values: Iterable[str], *, minimum: int, maximum: int) -> int:
    longest = max((len(value) for value in values), default=minimum)
    return max(minimum, min(maximum, longest + 2))


def _record_dict(record: WireRecord) -> dict[str, Any]:
    return record.model_dump(exclude_none=True)


def _first(data: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            return value
    return None


def _clean_cell(value: Any) -> Any:
    if value == "":
        return None
    return value


def _table_cell(value: Any) -> Any:
    if value == "":
        return None
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def _format_source_pages(value: Any) -> str | None:
    if not value:
        return None
    if isinstance(value, (list, tuple, set)):
        pages = sorted({int(page) for page in value if str(page).strip().isdigit()})
        return ",".join(str(page) for page in pages) or None
    return str(value)


def _format_string_list(value: Any) -> str | None:
    if not value:
        return None
    if isinstance(value, (list, tuple, set)):
        values = [str(item) for item in value if item not in (None, "")]
        return ",".join(dict.fromkeys(values)) or None
    return str(value)


def _format_json_cell(value: Any) -> str | None:
    if value in (None, "", [], {}):
        return None
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def records_by_segment_to_xlsx_bytes(
    records_by_segment: Mapping[str, Sequence[WireRecord]],
    *,
    terminal_strip_mapping: Mapping[str, str] | None = None,
    export_fields: ExportFields = PREPARED_EXPORT_FIELDS,
) -> bytes:
    """Create one workbook containing one wiring-table/raw pair per segment."""
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    used_titles: set[str] = set()
    for segment_id, records in records_by_segment.items():
        title = _unique_sheet_title(str(segment_id) or "wiring-table", used_titles)
        table_sheet = workbook.create_sheet(title)
        _write_wire_table(
            table_sheet, records, terminal_strip_mapping=terminal_strip_mapping, export_fields=export_fields
        )
        raw_sheet = workbook.create_sheet(_unique_sheet_title(f"{title}-raw", used_titles))
        _write_raw_table(raw_sheet, records)

    if not workbook.worksheets:
        sheet = workbook.create_sheet("wiring-table")
        _write_wire_table(
            sheet, [], terminal_strip_mapping=terminal_strip_mapping, export_fields=export_fields
        )

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _safe_sheet_title(title: str) -> str:
    invalid = set("[]:*?/\\")
    cleaned = "".join("_" if char in invalid else char for char in title).strip() or "放线表"
    return cleaned[:31]


def _unique_sheet_title(title: str, used_titles: set[str]) -> str:
    base = _safe_sheet_title(title)
    candidate = base
    suffix = 1
    while candidate in used_titles:
        marker = str(suffix)
        candidate = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    used_titles.add(candidate)
    return candidate
