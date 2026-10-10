from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from agent_service.domain.models.wiring import (
    WireRecord,
)
from agent_service.profiles.export_fields import PREPARED_EXPORT_FIELDS, ExportFields

"""Writer for the fixed downstream wiring-table layout.

The input template is legacy BIFF `.xls`.  `xlwt` is intentionally optional so
the API can report a clear configuration error instead of returning an
incompatible `.xlsx` with an `.xls` suffix.
"""


TEMPLATE_HEADERS = [
    "页码",
    "原理号",
    "电压等级",
    "起点代号",
    "起点描述",
    "起点端子",
    "终点代号",
    "终点描述",
    "终点端子",
    "电流",
    "备注",
    "色标",
]
TEMPLATE_COLUMN_COUNT = 23
DEFAULT_TEMPLATE_WIDTHS = [
    2267,
    2889,
    2486,
    2706,
    2486,
    2779,
    3730,
    5595,
    2486,
    4388,
    2962,
    2962,
    2340,
    18322,
    2340,
    2340,
    2340,
    2340,
    2340,
    2340,
    2340,
    2340,
    15286,
]


@dataclass(frozen=True)
class TemplateSpec:
    sheet_name: str
    headers: tuple[str, ...]
    column_widths: tuple[int, ...]


def load_template_spec(template_path: Path) -> TemplateSpec:
    """Read the real BIFF template's sheet, headers, column count and widths."""
    if not template_path.is_file():
        raise FileNotFoundError(template_path)

    try:
        import xlrd
    except ImportError:
        # Keep native export usable in an already provisioned environment; the
        # project requirements include xlrd so normal runs use the real spec.
        return _default_template_spec()

    try:
        workbook = xlrd.open_workbook(str(template_path), formatting_info=True)
    except Exception as exc:
        raise RuntimeError(f"无法读取 XLS 模板: {template_path}") from exc
    if not workbook.sheet_names():
        raise RuntimeError(f"XLS template has no worksheet: {template_path}")
    sheet = workbook.sheet_by_index(0)
    headers = tuple(str(value or "") for value in sheet.row_values(0))
    column_count = max(len(headers), TEMPLATE_COLUMN_COUNT)
    headers = headers + ("",) * (column_count - len(headers))
    widths = tuple(
        int(sheet.colinfo_map[index].width)
        if index in sheet.colinfo_map
        else DEFAULT_TEMPLATE_WIDTHS[index]
        if index < len(DEFAULT_TEMPLATE_WIDTHS)
        else 2340
        for index in range(column_count)
    )
    return TemplateSpec(
        sheet_name=sheet.name or "原理图导入表格",
        headers=headers,
        column_widths=widths,
    )


def _default_template_spec() -> TemplateSpec:
    headers = tuple(TEMPLATE_HEADERS) + ("",) * (TEMPLATE_COLUMN_COUNT - len(TEMPLATE_HEADERS))
    return TemplateSpec(
        sheet_name="原理图导入表格",
        headers=headers,
        column_widths=tuple(DEFAULT_TEMPLATE_WIDTHS),
    )


def records_to_template_xls_bytes(
    records: Sequence[WireRecord],
    *,
    template_path: Path,
    terminal_strip_mapping: dict[str, str] | None = None,
    export_fields: ExportFields = PREPARED_EXPORT_FIELDS,
) -> bytes:
    """Generate a native `.xls` workbook using the supplied template layout."""
    try:
        import xlwt
    except ImportError as exc:
        raise RuntimeError(
            "严格 XLS 导出需要 xlwt。请在 backend 环境安装 requirements.txt 中的 xlwt，"
            "或先使用 wiring-table.xlsx 预览导出。"
        ) from exc

    spec = load_template_spec(template_path)
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet(spec.sheet_name[:31])
    header_style = xlwt.easyxf(
        "font: bold on; align: horiz center, vert center; pattern: pattern solid, fore_colour gray25"
    )
    body_style = xlwt.easyxf("align: vert center; align: wrap on")
    for col, header in enumerate(spec.headers):
        sheet.write(0, col, header, header_style)
        sheet.col(col).width = spec.column_widths[col]
    for row, record in enumerate(records, start=1):
        for col, value in enumerate(_template_row(record, terminal_strip_mapping or {}, export_fields)):
            if value is not None:
                sheet.write(row, col, value, body_style)
    sheet.panes_frozen = True
    sheet.horz_split_pos = 1
    from io import BytesIO

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def records_to_template_xlsx_bytes(
    records: Sequence[WireRecord],
    *,
    template_path: Path,
    terminal_strip_mapping: dict[str, str] | None = None,
    export_fields: ExportFields = PREPARED_EXPORT_FIELDS,
) -> bytes:
    """Generate an OOXML workbook with the same columns as the XLS template."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    spec = load_template_spec(template_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.sheet_name[:31]
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for column, header in enumerate(spec.headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, record in enumerate(records, start=2):
        for column, value in enumerate(
            _template_row(record, terminal_strip_mapping or {}, export_fields), start=1
        ):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in enumerate(spec.column_widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = max(1, width / 256)
    sheet.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _template_row(
    record: WireRecord, mapping: Mapping[str, str], export_fields: ExportFields = PREPARED_EXPORT_FIELDS
) -> list[Any]:
    start_device = _text(record.start_device)
    end_device = _text(record.end_device)
    strip = export_fields.terminal_strip(start_device, record.terminal_strip, mapping)
    pages = ",".join(str(page) for page in sorted(set(record.source_pages))) or None
    remarks = _join_text(
        record.remark,
        record.source_type,
        f"PDF页:{pages}" if pages else None,
        "外部资料待补" if record.external_source_required else None,
    )
    return [
        _text(record.drawing_page) or pages,
        _text(record.line_number) or _text(record.wire_number),
        strip,
        _text(record.start_location) or start_device,
        _text(record.start_name) or _text(record.start_part),
        _terminal(start_device, record.start_terminal),
        _text(record.end_location) or end_device,
        _text(record.end_name) or _text(record.end_part),
        _terminal(end_device, record.end_terminal),
        _text(record.attribute) or _text(record.model),
        remarks,
        _text(record.color),
    ]


def _terminal(device: str | None, terminal: Any) -> str | None:
    if terminal in (None, ""):
        return None
    value = str(terminal).strip()
    if value in {"*", "PE"} or ":" in value:
        return value
    return f"{device.lstrip('-')}:{value}" if device else value


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _join_text(*values: Any) -> str | None:
    items = [text for value in values if (text := _text(value))]
    return "; ".join(items) or None
