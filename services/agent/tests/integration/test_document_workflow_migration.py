from __future__ import annotations

import ast
import asyncio
import json
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from agent_service.application.three_stage_extraction import run_three_stage_extraction
from agent_service.config import AgentSettings
from agent_service.domain.enums import CrossPageState
from agent_service.domain.models.extraction_stages import (
    CrossPageCompletionData,
    CrossPageCompletionRequest,
    DrawingEndpoint,
    DrawingReference,
    PageClassificationData,
    PageClassificationRequest,
    PageScanData,
    PageScanRequest,
    ScannedConnection,
    ScannedWireUnit,
)
from agent_service.profiles import ProfileRegistry
from agent_service.profiles.extraction_policy import default_extraction_policies

ROOT = Path(__file__).resolve().parents[4]
SOURCE = ROOT / "services/agent/src/agent_service"


class ConnectionStages:
    def __init__(self) -> None:
        self.classifications: list[int] = []
        self.scans: list[int] = []
        self.targets: list[int] = []

    async def classify_page(self, request: PageClassificationRequest) -> PageClassificationData:
        self.classifications.append(request.page.pdf_page_number)
        return PageClassificationData(
            plant_function="002.C",
            drawing_page_number=request.page.pdf_page_number,
            confidence=0.99,
        )

    async def scan_page(self, request: PageScanRequest) -> PageScanData:
        number = request.page.pdf_page_number
        self.scans.append(number)
        if number == 2:
            return PageScanData(pdf_page_number=number)
        return PageScanData(
            pdf_page_number=number,
            units=[
                ScannedWireUnit(
                    unit_id="source-unit",
                    wire_number="27",
                    connections=[
                        ScannedConnection(
                            start=DrawingEndpoint(device="-XD3", terminal="1", name="Supply"),
                            line_number="003G01",
                            current="400 A",
                            current_basis="Ir",
                            current_source_text="Ir=400A",
                            is_cross_page=CrossPageState.CROSS_PAGE,
                            references=[DrawingReference(raw="=002.C+M/2.1")],
                        ),
                        ScannedConnection(
                            start=DrawingEndpoint(device="-MOTOR", terminal="U1"),
                            end=DrawingEndpoint(device="-XD4", terminal="2"),
                            line_number="003G02",
                            is_cross_page=CrossPageState.SAME_PAGE,
                        ),
                        ScannedConnection(
                            start=DrawingEndpoint(device="-M1", terminal="U1"),
                            end=DrawingEndpoint(device="-M2", terminal="U2"),
                            line_number="REJECTED",
                            is_cross_page=CrossPageState.SAME_PAGE,
                        ),
                    ],
                )
            ],
        )

    async def resolve_cross_page(self, request: CrossPageCompletionRequest) -> CrossPageCompletionData:
        self.targets.append(request.target_page.pdf_page_number)
        assert request.target_page.pdf_page_number == 2
        context = json.loads(request.task_context)
        assert context["source_record"]["start_device"] == "-XD3"
        return CrossPageCompletionData(
            task_id=request.task_id,
            end=DrawingEndpoint(device="-MOTOR", terminal="U2"),
            status="resolved",
            confidence=0.95,
        )


def test_nonempty_pdf_preserves_fields_and_resumes_without_model_calls(tmp_path: Path) -> None:
    from agent_service.infrastructure.document.pdf_runtime import load_pdf_module

    pdf = tmp_path / "drawing.pdf"
    document = load_pdf_module().open()
    for number in (1, 2):
        page = document.new_page()
        page.insert_text((72, 72), f"Plant Function: 002.C\nPage Number: {number}")
    document.save(pdf)
    document.close()
    registry = ProfileRegistry(
        profile_root=ROOT / "services/agent/profiles",
        workspace_root=ROOT,
        allow_legacy_references=True,
    )
    registry.load_all()
    stages = ConnectionStages()

    async def run() -> None:
        kwargs = dict(
            pdf_path=pdf,
            output_dir=tmp_path / "run",
            profile=registry.bind("zh"),
            settings=AgentSettings(workspace_root=ROOT),
            run_id="migration-run",
            project_id="migration-project",
        )
        with patch(
            "agent_service.application.three_stage_extraction.build_extraction_stage_dispatcher",
            return_value=stages,
        ):
            first = await run_three_stage_extraction(**kwargs)
            records = [
                record for records in first.execution.state["wiring_records"].values() for record in records
            ]
            rows = {record.line_number: record for record in records}
            assert set(rows) == {"003G01", "003G02"}
            cross = rows["003G01"]
            assert (cross.start_device, cross.start_terminal, cross.end_device, cross.end_terminal) == (
                "-XD3",
                "XD3:1",
                "-MOTOR",
                "U2",
            )
            assert cross.current == "400A"
            assert cross.source_pages == [1, 2]
            assert cross.is_cross_page == CrossPageState.CROSS_PAGE
            assert (rows["003G02"].start_device, rows["003G02"].end_device) == ("-XD4", "-MOTOR")
            before = json.loads((tmp_path / "run/table.json").read_text(encoding="utf-8"))
            workbook = load_workbook(first.final_xlsx, read_only=True)
            sheet = workbook.active
            assert sheet is not None
            xlsx = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
            assert xlsx == [before["headers"], *before["rows"]]
            assert stages.classifications == [1, 2]
            assert stages.scans == [1, 2]
            assert stages.targets == [2]
            second = await run_three_stage_extraction(**kwargs)
            after = json.loads((tmp_path / "run/table.json").read_text(encoding="utf-8"))
            assert before == after
            assert stages.classifications == [1, 2]
            assert stages.scans == [1, 2]
            assert stages.targets == [2]
            assert second.execution.state["connection_records"] == first.execution.state["connection_records"]
            assert not list((tmp_path / "run/pages").rglob("*__pdf_*.png"))

    asyncio.run(run())


def test_production_sources_do_not_import_zh_workflow_or_vendor_rules_in_graph() -> None:
    for path in SOURCE.rglob("*.py"):
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if isinstance(node, ast.ImportFrom):
                assert "zh_workflow" not in (node.module or ""), str(path)
                if "graphs" in path.parts:
                    assert ".profiles.zh" not in (node.module or ""), str(path)
            elif isinstance(node, ast.Import):
                assert all("zh_workflow" not in alias.name for alias in node.names), str(path)


def test_unregistered_policy_does_not_inherit_zh_rules() -> None:
    import pytest

    with pytest.raises(ValueError, match="No deterministic extraction policy"):
        default_extraction_policies().resolve("abb_native")


def test_wiring_models_reuse_shared_domain_models() -> None:
    from agent_service.domain.models.wiring import WireConnection

    assert WireConnection is ScannedConnection


def test_export_writer_uses_injected_zh_fields() -> None:
    from agent_service.domain.models.wiring import WireRecord
    from agent_service.infrastructure.document.excel import records_to_xlsx_bytes
    from agent_service.profiles.zh.export_fields import ZhExportFields

    record = WireRecord(start_device="-XD3", start_terminal="1", line_number="003G01")
    prepared = records_to_xlsx_bytes([record], export_fields=ZhExportFields())
    workbook = load_workbook(BytesIO(prepared), read_only=True)
    sheet = workbook.active
    assert sheet is not None
    assert sheet.cell(2, 9).value == "X3"
    assert sheet.cell(2, 13).value == "XD3:1"
    assert sheet.cell(2, 8).value == "003G01"
    workbook.close()


def test_retired_batch_modules_and_exports_are_unavailable() -> None:
    import importlib.util

    import pytest

    from agent_service.graphs.document_extraction import workflow

    for module in (
        "agent_service.application.document_extraction.legacy",
        "agent_service.application.document_extraction.segment_decider",
    ):
        assert importlib.util.find_spec(module) is None
    for name in ("_segment_node", "_extract_wiring_node", "_assemble_xlsx_node"):
        with pytest.raises(AttributeError):
            getattr(workflow, name)


def test_zh_workflow_compatibility_package_is_removed() -> None:
    import importlib.util

    assert not (SOURCE / "zh_workflow").exists()
    assert importlib.util.find_spec("agent_service.zh_workflow") is None
