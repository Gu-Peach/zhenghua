from __future__ import annotations

import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from agent_service.application.three_stage_extraction import run_three_stage_extraction
from agent_service.config import AgentSettings
from agent_service.domain.models.extraction_stages import (
    CrossPageCompletionData,
    PageClassificationData,
    PageScanData,
)
from agent_service.profiles import ProfileRegistry

ROOT = Path(__file__).resolve().parents[4]


class FakeStages:
    async def classify_page(self, request):
        return PageClassificationData(
            plant_function="002.C",
            drawing_page_number=1,
            confidence=0.99,
            reason="offline workflow test",
        )

    async def scan_page(self, request):
        return PageScanData(
            pdf_page_number=request.page.pdf_page_number,
            drawing_function="002.C",
            drawing_page_number=1,
            units=[],
        )

    async def resolve_cross_page(self, request):
        return CrossPageCompletionData(task_id=request.task_id, status="needs_review", needs_review=True)


class StandaloneWorkflowTests(unittest.TestCase):
    def test_pdf_runs_through_local_graph_and_writes_stage_json_and_xlsx(self) -> None:
        async def run() -> None:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz

            registry = ProfileRegistry(
                profile_root=ROOT / "services" / "agent" / "profiles",
                workspace_root=ROOT,
                allow_legacy_references=True,
            )
            registry.load_all()
            profile = registry.bind("zh")
            settings = AgentSettings(workspace_root=ROOT)

            with tempfile.TemporaryDirectory() as temporary:
                temp = Path(temporary)
                pdf_path = temp / "drawing.pdf"
                document = fitz.open()
                page = document.new_page()
                page.insert_text((72, 72), "Plant Function: 002.C\nPage Number: 1")
                document.save(pdf_path)
                document.close()

                with patch(
                    "agent_service.application.three_stage_extraction.build_extraction_stage_dispatcher",
                    return_value=FakeStages(),
                ):
                    result = await run_three_stage_extraction(
                        pdf_path=pdf_path,
                        output_dir=temp / "run",
                        profile=profile,
                        settings=settings,
                        run_id="offline-run",
                        project_id="offline-project",
                        max_pdf_pages=0,
                    )

                self.assertTrue((result.stage_1_dir / "page-classifications.json").is_file())
                self.assertTrue((result.stage_1_dir / "drawing_index.json").is_file())
                self.assertTrue((result.stage_2_dir / "page-scan-results.json").is_file())
                self.assertTrue((result.stage_2_dir / "wire-units.json").is_file())
                self.assertTrue((result.stage_3_dir / "cross-page-tasks.json").is_file())
                self.assertTrue((result.stage_3_dir / "table.json").is_file())
                self.assertTrue(result.final_xlsx.is_file())
                self.assertTrue((temp / "run" / "pages" / "002.C" / "1.png").is_file())

                from openpyxl import load_workbook

                workbook = load_workbook(result.final_xlsx, read_only=True)
                self.assertEqual(workbook.active.max_column, 11)
                workbook.close()

        asyncio.run(run())


if __name__ == "__main__":
    unittest.main()
