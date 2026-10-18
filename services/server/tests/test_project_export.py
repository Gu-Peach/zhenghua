"""Offline API and OOXML checks; openpyxl is used only as an independent reader."""

from copy import deepcopy
from io import BytesIO
from uuid import uuid4

import pytest
import test_simple_management_api as export_fixtures
from openpyxl import load_workbook
from test_simple_management_api import PROJECT
from wiring_server.project_export import COLUMNS, XLSX_TYPE

management = export_fixtures.management


def test_export_all_columns_preserves_text_zero_and_empty(management):
    client, backend = management
    row = backend.tables["extraction_connection_rows"][0]
    row.update(
        principle_number="001002",
        terminal_strip="X22",
        start_code="+01F12.1",
        start_description="照明柜1",
        start_terminal="X3.2:101",
        end_code="ER-SO1",
        end_description="电气房单相维修插座#1",
        end_terminal="=003.C/1.1",
        current_value="0",
        remark='=HYPERLINK("https://invalid")',
        color_mark=None,
        evidence={"voltage_level": "230VAC"},
    )
    response = client.get(f"/api/v1/projects/{PROJECT}/export.xlsx")
    assert response.status_code == 200
    assert XLSX_TYPE in response.headers["content-type"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "private, no-store"
    book = load_workbook(BytesIO(response.content))
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == [col[0] for col in COLUMNS]
    assert sheet.max_column == 13 and sheet.max_row == 2
    assert [sheet.cell(2, i).value for i in (2, 3, 4, 7, 10, 11, 13)] == [
        "001002",
        "230VAC",
        "X22",
        "X3.2:101",
        "=003.C/1.1",
        "0",
        "",
    ]
    assert sheet["L2"].data_type == "s" and sheet["L2"].value.startswith("=HYPERLINK")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:M2"
    assert sheet["A1"].fill.fgColor.rgb == "FFFFFF00"
    assert all(request.method == "GET" for request in backend.calls)
    assert not backend.uploads


def test_export_paginates_entire_active_project_without_other_batch_or_owner(
    management,
):
    client, backend = management
    original = backend.tables["extraction_connection_rows"][0]
    rows = [
        {
            **deepcopy(original),
            "id": str(uuid4()),
            "row_order": i,
            "principle_number": f"P{i:04}",
        }
        for i in range(1002)
    ]
    backend.tables["extraction_projects"][0]["drawing_index"]["active_batch_id"] = (
        "run-one"
    )
    backend.tables["extraction_connection_rows"] = [
        *rows,
        {**original, "id": str(uuid4()), "batch_id": "old"},
        {**original, "id": str(uuid4()), "project_id": str(uuid4())},
    ]
    response = client.get(f"/api/v1/projects/{PROJECT}/export.xlsx")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), read_only=True)
    sheet = book.active
    assert sheet.max_row == 1003
    assert sheet["B1003"].value == "P1001"
    calls = [
        call
        for call in backend.calls
        if call.url.path.endswith("/extraction_connection_rows")
        and call.url.params.get("offset") == "1000"
    ]
    assert (
        len(calls) == 2
    )  # Data and final version verification both traverse all rows.


def test_export_empty_project_is_header_only(management):
    client, backend = management
    backend.tables["extraction_connection_rows"] = []
    result = client.get(f"/api/v1/projects/{PROJECT}/export.xlsx")
    assert result.status_code == 200
    sheet = load_workbook(BytesIO(result.content)).active
    assert sheet.max_row == 1 and sheet.max_column == 13


def test_export_checks_login_owner_and_processing(management):
    client, backend = management
    path = f"/api/v1/projects/{PROJECT}/export.xlsx"
    assert client.get(path, headers={"Authorization": ""}).status_code == 401
    assert (
        client.get(path, headers={"Authorization": "Bearer other"}).status_code == 403
    )
    backend.tables["extraction_projects"][0]["status"] = "processing"
    response = client.get(path)
    assert (
        response.status_code == 409
        and response.json()["detail"]["code"] == "EXTRACTION_IN_PROGRESS"
    )


@pytest.mark.parametrize("change", ["version", "delete", "insert", "project", "owner"])
def test_export_rechecks_versions_and_access_before_returning_file(
    management, monkeypatch, change
):
    from wiring_server import project_export

    client, backend = management
    builder = project_export.build_xlsx

    def changed(rows):
        content = builder(rows)
        stored = backend.tables["extraction_connection_rows"]
        if change == "version":
            stored[0]["record_version"] += 1
        elif change == "delete":
            stored.clear()
        elif change == "insert":
            stored.append({**stored[0], "id": str(uuid4())})
        elif change == "project":
            backend.tables["extraction_projects"][0]["status"] = "processing"
        else:
            backend.tables["extraction_projects"][0]["owner_id"] = "another-owner"
        return content

    monkeypatch.setattr(project_export, "build_xlsx", changed)
    result = client.get(f"/api/v1/projects/{PROJECT}/export.xlsx")
    assert result.status_code == (403 if change == "owner" else 409)
    assert "spreadsheetml" not in result.headers["content-type"]


def test_export_rejects_invalid_cells_without_losing_text(management):
    client, backend = management
    backend.tables["extraction_connection_rows"][0]["remark"] = "x" * 32768
    result = client.get(f"/api/v1/projects/{PROJECT}/export.xlsx")
    assert (
        result.status_code == 422
        and result.json()["detail"]["code"] == "EXPORT_CELL_INVALID"
    )


def test_export_contract_is_binary_and_does_not_accept_client_scope(management):
    client, _ = management
    operation = client.get("/openapi.json").json()["paths"][
        "/api/v1/projects/{project_id}/export.xlsx"
    ]["get"]
    assert operation["responses"]["200"]["content"][XLSX_TYPE]["schema"] == {
        "type": "string",
        "format": "binary",
    }
    assert not any(param["in"] == "query" for param in operation["parameters"])
    response = client.get(
        f"/api/v1/projects/{PROJECT}/export.xlsx",
        headers={"Origin": "http://localhost:5174"},
    )
    assert "Content-Disposition" in response.headers["access-control-expose-headers"]
